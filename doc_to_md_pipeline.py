#!/usr/bin/env python3
"""Batch-convert enterprise documents into Markdown for AnythingLLM."""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from tenacity import Retrying, retry_if_exception, stop_after_attempt, wait_exponential
from tqdm import tqdm


SUPPORTED_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt"}
LEGACY_TARGETS = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}
MANIFEST_FILENAME = ".conversion_manifest.json"
ERROR_LOG_FILENAME = "conversion_error.log"
DEFAULT_MODEL = "deepseek-flash"
DEFAULT_BASE_URL = "https://api.deepseek.com"
SYSTEM_PROMPT = (
    "你是一個專業的資料結構化工程師。請精確將圖片內容轉譯為 Markdown 格式。"
    "若為表格，請轉為對齊的 Markdown Table；若為電路圖、架構圖或流程圖，"
    "請條列出圖中所有文字標籤，並詳細描述各區塊的連接邏輯與技術規格。"
    "直接輸出 Markdown，不包含額外問候語。"
)
ENHANCE_SYSTEM_PROMPT = (
    "你是一個嚴謹的 Markdown 文件編輯器。請只整理輸入內容的結構與排版，不得摘要、刪除、"
    "翻譯、推測或新增任何資訊。可以統一標題層級、段落、清單與空白，但必須逐字保留所有"
    "數值、日期、百分比、型號、專有名詞、URL、電子郵件與技術規格。直接輸出整理後的 Markdown，"
    "不要加入說明、問候語或 Markdown code fence。"
)
DEFAULT_ENHANCE_CHUNK_CHARS = 24_000
_SOFFICE_LOCK = threading.Lock()


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def hash_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def yaml_string(value: str) -> str:
    """JSON double-quoted strings are also valid YAML scalars."""
    return json.dumps(value, ensure_ascii=False)


def markdown_table(rows: Sequence[Sequence[Any]], headers: Sequence[Any] | None = None) -> str:
    def clean(value: Any) -> str:
        if value is None:
            return ""
        return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")

    materialized = [[clean(cell) for cell in row] for row in rows]
    width = max([len(row) for row in materialized] + ([len(headers)] if headers else [0]))
    if width == 0:
        return ""
    normalized = [row + [""] * (width - len(row)) for row in materialized]
    header = [clean(cell) for cell in headers] if headers else [f"欄位 {i + 1}" for i in range(width)]
    header += [""] * (width - len(header))
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized)
    return "\n".join(lines)


@dataclass(frozen=True)
class ContentPart:
    kind: str
    content: str | bytes
    label: str = ""
    mime_type: str = "image/png"


@dataclass
class ExtractionResult:
    parts: list[ContentPart] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class FileJob:
    source: Path
    relative_source: Path
    output: Path
    output_relative: Path
    sha256: str
    mtime: float


@dataclass(frozen=True)
class SourceItem:
    source: Path
    relative_source: Path


@dataclass
class ScanPlan:
    total: int
    skipped: int
    jobs: list[FileJob]
    manifest_changed: bool = False


@dataclass
class ConversionResult:
    job: FileJob
    success: bool
    image_requests: int = 0
    enhancement_requests: int = 0
    error: str | None = None
    warnings: list[str] = field(default_factory=list)


class Manifest:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.entries: dict[str, dict[str, Any]] = {}

    def load(self) -> None:
        if not self.path.exists():
            return
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuntimeError(f"無法讀取 Manifest：{self.path}: {exc}") from exc
        if not isinstance(data, dict):
            raise RuntimeError(f"Manifest 根節點必須是 JSON object：{self.path}")
        self.entries = data

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(prefix=f".{self.path.name}.", suffix=".tmp", dir=self.path.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
                json.dump(self.entries, stream, ensure_ascii=False, indent=2, sort_keys=True)
                stream.write("\n")
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temp_name, self.path)
        except Exception:
            try:
                os.unlink(temp_name)
            except OSError:
                pass
            raise


def scan_supported_files(input_dir: Path) -> list[Path]:
    return sorted(
        (path for path in input_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS),
        key=lambda path: path.relative_to(input_dir).as_posix().casefold(),
    )


def output_suffix(output_format: str) -> str:
    return {"md": ".md", "txt": ".txt", "json": ".json"}[output_format]


def build_output_mapping(
    files: Sequence[Path], input_dir: Path, output_dir: Path, suffix: str = ".md"
) -> dict[Path, tuple[Path, Path]]:
    """Map source paths to Markdown paths, disambiguating equal stems."""
    candidates: dict[str, list[Path]] = {}
    for source in files:
        relative = source.relative_to(input_dir)
        candidate = relative.with_suffix(suffix)
        candidates.setdefault(candidate.as_posix().casefold(), []).append(source)

    mapping: dict[Path, tuple[Path, Path]] = {}
    for grouped_sources in candidates.values():
        collision = len(grouped_sources) > 1
        for source in grouped_sources:
            relative = source.relative_to(input_dir)
            if collision:
                output_relative = relative.with_name(f"{relative.stem}{relative.suffix.lower()}{suffix}")
            else:
                output_relative = relative.with_suffix(suffix)
            mapping[source] = (output_dir / output_relative, output_relative)
    return mapping


def build_item_output_mapping(
    items: Sequence[SourceItem], output_dir: Path, suffix: str
) -> dict[Path, tuple[Path, Path]]:
    candidates: dict[str, list[SourceItem]] = {}
    for item in items:
        candidate = item.relative_source.with_suffix(suffix)
        candidates.setdefault(candidate.as_posix().casefold(), []).append(item)

    mapping: dict[Path, tuple[Path, Path]] = {}
    for grouped_items in candidates.values():
        collision = len(grouped_items) > 1
        for item in grouped_items:
            relative = item.relative_source
            if collision:
                output_relative = relative.with_name(f"{relative.stem}{item.source.suffix.lower()}{suffix}")
            else:
                output_relative = relative.with_suffix(suffix)
            mapping[item.source] = (output_dir / output_relative, output_relative)
    return mapping


def _unique_anchor(name: str, used: set[str]) -> str:
    base = name or "input"
    candidate = base
    index = 2
    while candidate.casefold() in used:
        candidate = f"{base}-{index}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def collect_source_items(
    input_paths: Sequence[Path],
    allowed_extensions: set[str] | None = None,
    excluded_paths: Sequence[Path] | None = None,
) -> list[SourceItem]:
    """Expand one or more dropped files/directories into stable relative paths."""
    allowed = {value.lower() for value in (allowed_extensions or SUPPORTED_EXTENSIONS)}
    paths = [path.expanduser().resolve() for path in input_paths]
    exclusions = [path.expanduser().resolve() for path in (excluded_paths or [])]
    if not paths:
        raise ValueError("請至少選擇一個輸入檔案或資料夾")
    missing = [path for path in paths if not path.exists()]
    if missing:
        raise ValueError(f"輸入路徑不存在：{missing[0]}")

    expanded: list[tuple[Path, Path]] = []
    if len(paths) == 1 and paths[0].is_dir():
        root = paths[0]
        expanded.extend((source, source.relative_to(root)) for source in scan_supported_files(root))
    else:
        try:
            bases = [path if path.is_dir() else path.parent for path in paths]
            common_root = Path(os.path.commonpath([str(path) for path in bases]))
            same_drive = True
        except ValueError:
            common_root = Path()
            same_drive = False

        used_anchors: set[str] = set()
        for selected in paths:
            if selected.is_file():
                if selected.suffix.lower() in allowed:
                    if same_drive:
                        relative = selected.relative_to(common_root)
                    else:
                        anchor = _unique_anchor(selected.parent.name, used_anchors)
                        relative = Path(anchor) / selected.name
                    expanded.append((selected, relative))
                continue

            anchor = _unique_anchor(selected.name, used_anchors)
            for source in scan_supported_files(selected):
                if source.suffix.lower() not in allowed:
                    continue
                relative = source.relative_to(common_root) if same_drive else Path(anchor) / source.relative_to(selected)
                expanded.append((source, relative))

    deduplicated: dict[str, SourceItem] = {}
    for source, relative in expanded:
        if source.suffix.lower() not in allowed:
            continue
        if any(source == excluded or source.is_relative_to(excluded) for excluded in exclusions):
            continue
        deduplicated.setdefault(str(source).casefold(), SourceItem(source, relative))
    return sorted(deduplicated.values(), key=lambda item: item.relative_source.as_posix().casefold())


def create_scan_plan_from_items(
    items: Sequence[SourceItem],
    output_dir: Path,
    manifest: Manifest,
    force: bool,
    output_format: str = "md",
    model: str | None = None,
    base_url: str | None = None,
    processing_mode: str = "hybrid",
) -> ScanPlan:
    outputs = build_item_output_mapping(items, output_dir, output_suffix(output_format))
    jobs: list[FileJob] = []
    skipped = 0
    changed = False

    for item in items:
        source = item.source
        relative = item.relative_source
        key = relative.as_posix()
        output, output_relative = outputs[source]
        stat = source.stat()
        entry = manifest.entries.get(key)
        expected_output = output_relative.as_posix()
        settings_match = not isinstance(entry, dict) or (
            (model is None or entry.get("model") == model)
            and (base_url is None or entry.get("base_url") == base_url)
            and entry.get("output_format", "md") == output_format
            and entry.get("processing_mode", "hybrid") == processing_mode
        )

        can_fast_skip = (
            not force
            and isinstance(entry, dict)
            and settings_match
            and entry.get("status") == "success"
            and entry.get("last_modified") == stat.st_mtime
            and entry.get("output_md") == expected_output
            and output.is_file()
        )
        if can_fast_skip:
            skipped += 1
            continue

        sha256 = hash_file(source)
        can_hash_skip = (
            not force
            and isinstance(entry, dict)
            and settings_match
            and entry.get("status") == "success"
            and entry.get("sha256") == sha256
            and entry.get("output_md") == expected_output
            and output.is_file()
        )
        if can_hash_skip:
            entry["last_modified"] = stat.st_mtime
            skipped += 1
            changed = True
            continue

        jobs.append(FileJob(source, relative, output, output_relative, sha256, stat.st_mtime))
    return ScanPlan(total=len(items), skipped=skipped, jobs=jobs, manifest_changed=changed)


def create_scan_plan(
    input_dir: Path,
    output_dir: Path,
    manifest: Manifest,
    force: bool,
    output_format: str = "md",
    processing_mode: str = "hybrid",
) -> ScanPlan:
    files = scan_supported_files(input_dir)
    items = [SourceItem(source, source.relative_to(input_dir)) for source in files]
    return create_scan_plan_from_items(
        items, output_dir, manifest, force, output_format, processing_mode=processing_mode
    )


def normalize_image(data: bytes) -> bytes:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        image.seek(0)
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGB")
        output = io.BytesIO()
        image.save(output, format="PNG", optimize=True)
        return output.getvalue()


def text_part(text: str, label: str = "") -> ContentPart | None:
    stripped = text.strip()
    return ContentPart("text", stripped, label) if stripped else None


def image_part(data: bytes, label: str) -> ContentPart:
    return ContentPart("image", normalize_image(data), label, "image/png")


def read_text_file(path: Path) -> ExtractionResult:
    encodings = ("utf-8-sig", "utf-16", "cp950", "cp1252")
    for encoding in encodings:
        try:
            return ExtractionResult(parts=[ContentPart("text", path.read_text(encoding=encoding))])
        except UnicodeDecodeError:
            continue
    return ExtractionResult(parts=[ContentPart("text", path.read_text(encoding="utf-8", errors="replace"))])


def extract_xlsx(path: Path) -> ExtractionResult:
    import pandas as pd
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    parts: list[ContentPart] = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            while rows and all(value is None for value in rows[-1]):
                rows.pop()
            if rows:
                width = max(len(row) for row in rows)
                while width > 0 and all((row[width - 1] if len(row) >= width else None) is None for row in rows):
                    width -= 1
                rows = [tuple(row[:width]) for row in rows]
            parts.append(ContentPart("text", f"## 工作表：{sheet.title}"))
            if not rows or not rows[0]:
                parts.append(ContentPart("text", "_空白工作表_"))
                continue
            dataframe = pd.DataFrame(rows).fillna("")
            headers = [f"欄位 {index + 1}" for index in range(len(dataframe.columns))]
            parts.append(ContentPart("text", dataframe.to_markdown(index=False, headers=headers)))
    finally:
        workbook.close()
    return ExtractionResult(parts=parts)


def extract_docx(path: Path) -> ExtractionResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    document = Document(path)
    parts: list[ContentPart] = []
    warnings: list[str] = []
    image_index = 0

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                style_name = (paragraph.style.name or "") if paragraph.style else ""
                if style_name.lower().startswith("heading"):
                    try:
                        level = min(6, max(1, int(style_name.split()[-1])))
                    except (ValueError, IndexError):
                        level = 2
                    parts.append(ContentPart("text", f"{'#' * level} {paragraph.text.strip()}"))
                else:
                    parts.append(ContentPart("text", paragraph.text.strip()))

            for blip in child.xpath(".//a:blip"):
                relation_id = blip.get(qn("r:embed"))
                if not relation_id:
                    continue
                try:
                    blob = document.part.related_parts[relation_id].blob
                    image_index += 1
                    parts.append(image_part(blob, f"內嵌圖片 {image_index}"))
                except Exception as exc:
                    warnings.append(f"略過無法解碼的 DOCX 圖片：{exc}")
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                parts.append(ContentPart("text", markdown_table(rows[1:], headers=rows[0])))

    return ExtractionResult(parts=parts, warnings=warnings)


def _iter_ppt_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_ppt_shapes(shape.shapes)
        else:
            yield shape


def _extract_chart_markdown(chart: Any) -> str:
    rows: list[list[Any]] = []
    series_list = list(chart.series)
    if not series_list:
        return ""
    max_values = max((len(list(series.values)) for series in series_list), default=0)
    for index in range(max_values):
        row: list[Any] = [index + 1]
        for series in series_list:
            values = list(series.values)
            row.append(values[index] if index < len(values) else "")
        rows.append(row)
    headers = ["項目"] + [series.name or f"數列 {index + 1}" for index, series in enumerate(series_list)]
    return markdown_table(rows, headers=headers)


def find_soffice() -> str | None:
    for executable in ("soffice", "libreoffice"):
        if found := shutil.which(executable):
            return found
    if os.name == "nt":
        candidates = (
            Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "LibreOffice/program/soffice.exe",
            Path(os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)")) / "LibreOffice/program/soffice.exe",
        )
        for candidate in candidates:
            if candidate.is_file():
                return str(candidate)
    return None


def libreoffice_convert(source: Path, target_extension: str, output_dir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("找不到 LibreOffice soffice；請安裝 LibreOffice 並加入 PATH")
    output_dir.mkdir(parents=True, exist_ok=True)
    target_format = target_extension.lstrip(".")
    run_options: dict[str, Any] = {
        "capture_output": True,
        "text": True,
        "timeout": 300,
        "check": False,
    }
    if os.name == "nt":
        run_options["creationflags"] = subprocess.CREATE_NO_WINDOW
    with _SOFFICE_LOCK:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", target_format, "--outdir", str(output_dir), str(source)],
            **run_options,
        )
    expected = output_dir / f"{source.stem}.{target_format}"
    if result.returncode != 0 or not expected.is_file():
        detail = (result.stderr or result.stdout or "未產生輸出檔").strip()
        raise RuntimeError(f"LibreOffice 轉檔失敗 ({source.name} -> {target_format})：{detail}")
    return expected


def render_presentation_slides(path: Path) -> tuple[dict[int, bytes], str | None]:
    if not find_soffice():
        return {}, "未偵測到 LibreOffice，無法產生完整投影片預覽；仍會處理文字、表格與內嵌圖片"
    try:
        import pymupdf as fitz

        with tempfile.TemporaryDirectory(prefix="doc-to-md-ppt-") as temp_name:
            pdf_path = libreoffice_convert(path, ".pdf", Path(temp_name))
            document = fitz.open(pdf_path)
            try:
                slides = {
                    index + 1: page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
                    for index, page in enumerate(document)
                }
            finally:
                document.close()
            return slides, None
    except Exception as exc:
        return {}, f"無法產生完整投影片預覽：{exc}"


def extract_pptx(path: Path) -> ExtractionResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    rendered_slides, render_warning = render_presentation_slides(path)
    warnings = [render_warning] if render_warning else []
    parts: list[ContentPart] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts.append(ContentPart("text", f"## 投影片 {slide_index}"))
        has_visual_structure = False
        for shape in _iter_ppt_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(ContentPart("text", shape.text.strip()))
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    parts.append(ContentPart("text", markdown_table(rows[1:], headers=rows[0])))
            if getattr(shape, "has_chart", False):
                has_visual_structure = True
                chart_markdown = _extract_chart_markdown(shape.chart)
                if chart_markdown:
                    parts.append(ContentPart("text", f"### 圖表資料\n\n{chart_markdown}"))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    parts.append(image_part(shape.image.blob, f"投影片 {slide_index} 內嵌圖片"))
                except Exception as exc:
                    warnings.append(f"略過投影片 {slide_index} 無法解碼的圖片：{exc}")
            elif shape.shape_type not in (MSO_SHAPE_TYPE.PLACEHOLDER,):
                has_visual_structure = True

        if has_visual_structure and slide_index in rendered_slides:
            parts.append(image_part(rendered_slides[slide_index], f"投影片 {slide_index} 完整視圖"))

    return ExtractionResult(parts=parts, warnings=warnings)


def extract_pdf(path: Path) -> ExtractionResult:
    import pymupdf as fitz

    document = fitz.open(path)
    parts: list[ContentPart] = []
    warnings: list[str] = []
    try:
        for page_index, page in enumerate(document, start=1):
            parts.append(ContentPart("text", f"## 第 {page_index} 頁"))
            page_dict = page.get_text("dict", sort=True)
            page_text: list[str] = []
            image_index = 0
            for block in page_dict.get("blocks", []):
                if block.get("type") == 0:
                    lines = []
                    for line in block.get("lines", []):
                        line_text = "".join(span.get("text", "") for span in line.get("spans", []))
                        if line_text.strip():
                            lines.append(line_text)
                    if lines:
                        value = "\n".join(lines)
                        page_text.append(value)
                        parts.append(ContentPart("text", value))
                elif block.get("type") == 1 and block.get("image"):
                    try:
                        image_index += 1
                        parts.append(image_part(block["image"], f"第 {page_index} 頁圖片 {image_index}"))
                    except Exception as exc:
                        warnings.append(f"略過第 {page_index} 頁無法解碼的圖片：{exc}")

            table_count = 0
            try:
                finder = page.find_tables()
                for table_count, table in enumerate(finder.tables, start=1):
                    clip = fitz.Rect(table.bbox)
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=clip, alpha=False)
                    parts.append(image_part(pixmap.tobytes("png"), f"第 {page_index} 頁表格 {table_count}"))
            except Exception as exc:
                warnings.append(f"第 {page_index} 頁表格偵測失敗：{exc}")

            joined_text = "\n".join(page_text).strip()
            drawings = page.get_drawings()
            needs_page_image = not joined_text or (len(joined_text) < 1000 and len(drawings) >= 3 and table_count == 0)
            if needs_page_image and image_index == 0:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                parts.append(image_part(pixmap.tobytes("png"), f"第 {page_index} 頁完整視圖"))
    finally:
        document.close()
    return ExtractionResult(parts=parts, warnings=warnings)


def extract_document(path: Path) -> ExtractionResult:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return read_text_file(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix in LEGACY_TARGETS:
        with tempfile.TemporaryDirectory(prefix="doc-to-md-legacy-") as temp_name:
            converted = libreoffice_convert(path, LEGACY_TARGETS[suffix], Path(temp_name))
            return extract_document(converted)
    raise ValueError(f"不支援的檔案格式：{path.suffix}")


def is_retryable_api_error(exc: BaseException) -> bool:
    status_code = getattr(exc, "status_code", None)
    if status_code == 429 or (isinstance(status_code, int) and status_code >= 500):
        return True
    return exc.__class__.__name__ in {"APIConnectionError", "APITimeoutError", "RateLimitError", "InternalServerError"}


class VisionClient:
    def __init__(self, api_key: str, base_url: str, model: str) -> None:
        from openai import OpenAI

        self.client = OpenAI(api_key=api_key, base_url=base_url, timeout=120.0)
        self.model = model

    def _complete(self, messages: list[dict[str, Any]]) -> str:
        retrying = Retrying(
            stop=stop_after_attempt(5),
            wait=wait_exponential(multiplier=1, min=1, max=30),
            retry=retry_if_exception(is_retryable_api_error),
            reraise=True,
        )
        response = retrying(
            self.client.chat.completions.create,
            model=self.model,
            messages=messages,
            temperature=0,
        )
        content = response.choices[0].message.content
        if not content:
            raise RuntimeError("DeepSeek API 回傳空白內容")
        return content.strip()

    def describe(self, image: bytes, label: str) -> str:
        data_url = f"data:image/png;base64,{base64.b64encode(image).decode('ascii')}"
        return self._complete(
            [
                {"role": "system", "content": SYSTEM_PROMPT},
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": f"請轉譯此圖片（{label}）。"},
                        {"type": "image_url", "image_url": {"url": data_url, "detail": "original"}},
                    ],
                },
            ]
        )

    def enhance_markdown(self, markdown: str, chunk_index: int, chunk_count: int) -> str:
        return self._complete(
            [
                {"role": "system", "content": ENHANCE_SYSTEM_PROMPT},
                {
                    "role": "user",
                    "content": (
                        f"以下是文件 Markdown 的第 {chunk_index}/{chunk_count} 段。"
                        "請依規則整理，並完整保留原始資訊：\n\n"
                        f"{markdown}"
                    ),
                },
            ]
        )


def list_available_models(api_key: str, base_url: str) -> list[str]:
    """Query an OpenAI-compatible Models API without exposing the credential."""
    if not api_key:
        raise ValueError("查詢模型前請先輸入 API Key")
    from openai import OpenAI

    client = OpenAI(api_key=api_key, base_url=base_url.rstrip("/"), timeout=30.0)
    response = client.models.list()
    model_ids = sorted(
        {str(model.id).strip() for model in response.data if getattr(model, "id", None)},
        key=str.casefold,
    )
    if not model_ids:
        raise RuntimeError("端點回傳的模型清單為空")
    return model_ids


class VisionClientFactory:
    def __init__(self, api_key: str | None, base_url: str, model: str) -> None:
        self.api_key = api_key
        self.base_url = base_url
        self.model = model
        self.local = threading.local()

    def get(self) -> VisionClient:
        if not self.api_key:
            raise RuntimeError("需要 AI 處理，但未設定 DEEPSEEK_API_KEY")
        if not hasattr(self.local, "client"):
            self.local.client = VisionClient(self.api_key, self.base_url, self.model)
        return self.local.client


def split_markdown_frontmatter(markdown: str) -> tuple[str, str]:
    marker = "\n---\n\n"
    if markdown.startswith("---\n") and marker in markdown[4:]:
        boundary = markdown.find(marker, 4) + len(marker)
        return markdown[:boundary], markdown[boundary:]
    return "", markdown


def _is_protected_markdown_block(block: str) -> bool:
    stripped = block.strip()
    if stripped.startswith("```") or stripped.startswith("~~~"):
        return True
    lines = [line for line in stripped.splitlines() if line.strip()]
    if len(lines) < 2:
        return False
    table_lines = sum(line.lstrip().startswith("|") and line.rstrip().endswith("|") for line in lines)
    return table_lines / len(lines) >= 0.6


def _split_large_text(text: str, max_chars: int) -> list[str]:
    chunks: list[str] = []
    current: list[str] = []
    current_length = 0
    for line in text.splitlines():
        if len(line) > max_chars:
            if current:
                chunks.append("\n".join(current))
                current = []
                current_length = 0
            chunks.extend(line[index : index + max_chars] for index in range(0, len(line), max_chars))
            continue
        additional = len(line) + (1 if current else 0)
        if current and current_length + additional > max_chars:
            chunks.append("\n".join(current))
            current = []
            current_length = 0
        current.append(line)
        current_length += additional
    if current:
        chunks.append("\n".join(current))
    return chunks


def _split_markdown_blocks(body: str) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    fence_marker: str | None = None
    for line in body.strip().splitlines():
        stripped = line.lstrip()
        if fence_marker:
            current.append(line)
            if stripped.startswith(fence_marker):
                fence_marker = None
            continue
        if stripped.startswith("```") or stripped.startswith("~~~"):
            fence_marker = stripped[:3]
            current.append(line)
            continue
        if not line.strip():
            if current:
                blocks.append("\n".join(current).strip())
                current = []
            continue
        current.append(line)
    if current:
        blocks.append("\n".join(current).strip())
    return blocks


def split_markdown_for_enhancement(
    body: str, max_chars: int = DEFAULT_ENHANCE_CHUNK_CHARS
) -> list[tuple[str, bool]]:
    """Return ordered (content, should_enhance) segments while protecting tables and code."""
    if max_chars < 1000:
        raise ValueError("AI 整理分段大小不得小於 1000 字元")
    blocks = _split_markdown_blocks(body)
    segments: list[tuple[str, bool]] = []
    buffer: list[str] = []
    buffer_length = 0

    def flush_buffer() -> None:
        nonlocal buffer, buffer_length
        if buffer:
            segments.append(("\n\n".join(buffer), True))
            buffer = []
            buffer_length = 0

    protected_flags = [_is_protected_markdown_block(block) for block in blocks]
    for index, block in enumerate(blocks):
        heading_for_protected_block = (
            bool(re.fullmatch(r"#{1,6}\s+[^\n]+", block))
            and index + 1 < len(blocks)
            and protected_flags[index + 1]
        )
        if protected_flags[index] or heading_for_protected_block:
            flush_buffer()
            segments.append((block, False))
            continue
        if len(block) > max_chars:
            flush_buffer()
            segments.extend((chunk, True) for chunk in _split_large_text(block, max_chars) if chunk)
            continue
        additional = len(block) + (2 if buffer else 0)
        if buffer and buffer_length + additional > max_chars:
            flush_buffer()
        buffer.append(block)
        buffer_length += additional
    flush_buffer()
    return segments


_CRITICAL_TOKEN_PATTERN = re.compile(
    r"https?://[^\s<>()]+|"
    r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|"
    r"\b(?:[A-Za-z]+[A-Za-z0-9_-]*\d[A-Za-z0-9_-]*|\d+[A-Za-z][A-Za-z0-9_-]*)\b|"
    r"(?<![\w])(?:0x[0-9A-Fa-f]+|\d[\d,]*(?:\.\d+)?%?)(?![\w])"
)
_CONTENT_TOKEN_PATTERN = re.compile(r"[\u3400-\u9fff]|[A-Za-z][A-Za-z'_-]*")


def _critical_tokens(text: str) -> Counter[str]:
    tokens: list[str] = []
    for match in _CRITICAL_TOKEN_PATTERN.finditer(text):
        token = match.group(0).rstrip(".,;:!?")
        if re.fullmatch(r"\d[\d,]*(?:\.\d+)?%?", token):
            token = token.replace(",", "")
        tokens.append(token)
    return Counter(tokens)


def validate_enhanced_markdown(source: str, enhanced: str) -> tuple[bool, str]:
    if not enhanced.strip():
        return False, "模型回傳空白內容"
    if _critical_tokens(source) != _critical_tokens(enhanced):
        return False, "數值、網址、電子郵件或識別碼不一致"
    source_tokens = Counter(token.lower() for token in _CONTENT_TOKEN_PATTERN.findall(source))
    enhanced_tokens = Counter(token.lower() for token in _CONTENT_TOKEN_PATTERN.findall(enhanced))
    if source_tokens:
        retained = sum((source_tokens & enhanced_tokens).values()) / sum(source_tokens.values())
        added = sum((enhanced_tokens - source_tokens).values())
        if retained < 0.9 or added > max(5, int(sum(source_tokens.values()) * 0.15)):
            return False, "文字內容差異超出安全範圍"
    source_length = len(re.sub(r"\s+", "", source))
    enhanced_length = len(re.sub(r"\s+", "", enhanced))
    if source_length and not (0.65 <= enhanced_length / source_length <= 1.5):
        return False, "整理後內容長度變化超出安全範圍"
    return True, ""


def _strip_markdown_fence(value: str) -> str:
    stripped = value.strip()
    lines = stripped.splitlines()
    if len(lines) >= 2 and lines[0].strip().lower() in {"```", "```markdown", "```md"} and lines[-1].strip() == "```":
        return "\n".join(lines[1:-1]).strip()
    return stripped


def enhance_markdown_document(
    markdown: str,
    vision_factory: VisionClientFactory,
    max_chars: int = DEFAULT_ENHANCE_CHUNK_CHARS,
) -> tuple[str, int, list[str]]:
    frontmatter, body = split_markdown_frontmatter(markdown)
    segments = split_markdown_for_enhancement(body, max_chars)
    eligible_count = sum(should_enhance for _, should_enhance in segments)
    if eligible_count == 0:
        return markdown, 0, []

    client = vision_factory.get()
    enhanced_segments: list[str] = []
    warnings: list[str] = []
    request_count = 0
    eligible_index = 0
    for content, should_enhance in segments:
        if not should_enhance:
            enhanced_segments.append(content)
            continue
        eligible_index += 1
        request_count += 1
        try:
            candidate = _strip_markdown_fence(
                client.enhance_markdown(content, eligible_index, eligible_count)
            )
            valid, reason = validate_enhanced_markdown(content, candidate)
            if not valid:
                warnings.append(f"AI 整理第 {eligible_index}/{eligible_count} 段未通過驗證，保留原文：{reason}")
                candidate = content
        except Exception as exc:
            warnings.append(
                f"AI 整理第 {eligible_index}/{eligible_count} 段失敗，保留原文：{type(exc).__name__}: {exc}"
            )
            candidate = content
        enhanced_segments.append(candidate)

    enhanced_body = "\n\n".join(enhanced_segments).rstrip() + "\n"
    return frontmatter + enhanced_body, request_count, warnings


def assemble_markdown(
    relative_source: Path,
    model: str,
    processing_mode: str,
    converted_at: str,
    extraction: ExtractionResult,
    vision_factory: VisionClientFactory,
) -> tuple[str, int]:
    body: list[str] = []
    image_requests = 0
    for part in extraction.parts:
        if part.kind == "text":
            value = str(part.content).strip()
            if value:
                body.append(value)
        elif part.kind == "image":
            image_requests += 1
            description = vision_factory.get().describe(bytes(part.content), part.label)
            body.append(f"### 視覺內容：{part.label}\n\n{description}")
        else:
            raise ValueError(f"未知的內容區塊類型：{part.kind}")

    if not body:
        body.append("_無可擷取內容_ ")
    frontmatter = (
        "---\n"
        f"original_file: {yaml_string(relative_source.as_posix())}\n"
        f"converted_date: {yaml_string(converted_at)}\n"
        f"model: {yaml_string(model)}\n"
        f"processing_mode: {yaml_string(processing_mode)}\n"
        "---"
    )
    return frontmatter + "\n\n" + "\n\n".join(body).rstrip() + "\n", image_requests


def serialize_output(
    markdown: str,
    output_format: str,
    relative_source: Path,
    model: str,
    converted_at: str,
    processing_mode: str = "hybrid",
) -> str:
    marker = "\n---\n\n"
    body = markdown.split(marker, 1)[1] if marker in markdown else markdown
    if output_format == "md":
        return markdown
    if output_format == "txt":
        return (
            f"Original file: {relative_source.as_posix()}\n"
            f"Converted date: {converted_at}\n"
            f"Model: {model}\n"
            f"Processing mode: {processing_mode}\n\n"
            f"{body}"
        )
    if output_format == "json":
        return json.dumps(
            {
                "original_file": relative_source.as_posix(),
                "converted_date": converted_at,
                "model": model,
                "processing_mode": processing_mode,
                "content_markdown": body.rstrip(),
            },
            ensure_ascii=False,
            indent=2,
        ) + "\n"
    raise ValueError(f"不支援的輸出格式：{output_format}")


def atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except OSError:
            pass
        raise


def convert_job(
    job: FileJob,
    model: str,
    output_format: str,
    processing_mode: str,
    vision_factory: VisionClientFactory,
) -> ConversionResult:
    image_requests = 0
    enhancement_requests = 0
    try:
        extraction = extract_document(job.source)
        image_requests = sum(part.kind == "image" for part in extraction.parts)
        converted_at = utc_now()
        markdown, _ = assemble_markdown(
            job.relative_source, model, processing_mode, converted_at, extraction, vision_factory
        )
        warnings = list(extraction.warnings)
        if processing_mode == "ai-enhanced":
            markdown, enhancement_requests, enhancement_warnings = enhance_markdown_document(
                markdown, vision_factory
            )
            warnings.extend(enhancement_warnings)
        output = serialize_output(
            markdown,
            output_format,
            job.relative_source,
            model,
            converted_at,
            processing_mode,
        )
        atomic_write_text(job.output, output)
        return ConversionResult(
            job,
            True,
            image_requests=image_requests,
            enhancement_requests=enhancement_requests,
            warnings=warnings,
        )
    except Exception as exc:
        return ConversionResult(
            job,
            False,
            image_requests=image_requests,
            enhancement_requests=enhancement_requests,
            error=f"{type(exc).__name__}: {exc}",
        )


def configure_logger(output_dir: Path) -> logging.Logger:
    output_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("doc_to_md_pipeline")
    for existing_handler in list(logger.handlers):
        existing_handler.close()
        logger.removeHandler(existing_handler)
    logger.setLevel(logging.INFO)
    handler = logging.FileHandler(output_dir / ERROR_LOG_FILENAME, encoding="utf-8", delay=True)
    handler.setFormatter(logging.Formatter("%(asctime)s\t%(levelname)s\t%(message)s"))
    logger.addHandler(handler)
    logger.propagate = False
    return logger


def close_logger(logger: logging.Logger) -> None:
    for handler in list(logger.handlers):
        handler.close()
        logger.removeHandler(handler)


def update_manifest(
    manifest: Manifest,
    result: ConversionResult,
    model: str,
    base_url: str | None = None,
    output_format: str = "md",
    processing_mode: str = "hybrid",
) -> None:
    key = result.job.relative_source.as_posix()
    entry: dict[str, Any] = {
        "sha256": result.job.sha256,
        "last_modified": result.job.mtime,
        "output_md": result.job.output_relative.as_posix(),
        "converted_at": utc_now(),
        "status": "success" if result.success else "failed",
        "model": model,
        "output_format": output_format,
        "processing_mode": processing_mode,
    }
    if base_url:
        entry["base_url"] = base_url
    if result.error:
        entry["error"] = result.error
    manifest.entries[key] = entry


def validate_directories(input_dir: Path, output_dir: Path) -> tuple[Path, Path]:
    input_dir = input_dir.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    if not input_dir.is_dir():
        raise ValueError(f"輸入目錄不存在或不是目錄：{input_dir}")
    if input_dir == output_dir:
        raise ValueError("輸入與輸出目錄不可相同")
    return input_dir, output_dir


def run_pipeline_paths(
    args: argparse.Namespace,
    input_paths: Sequence[Path],
    *,
    allowed_extensions: set[str] | None = None,
    progress_callback: Any | None = None,
    quiet: bool = False,
) -> int:
    output_dir = Path(args.output_dir).expanduser().resolve()
    if any(path.expanduser().resolve() == output_dir for path in input_paths):
        raise ValueError("輸入與輸出目錄不可相同")
    output_dir.mkdir(parents=True, exist_ok=True)
    logger = configure_logger(output_dir)
    manifest = Manifest(output_dir / MANIFEST_FILENAME)
    manifest.load()

    output_format = getattr(args, "output_format", "md")
    processing_mode = getattr(args, "processing_mode", "hybrid")
    if processing_mode not in {"hybrid", "ai-enhanced"}:
        raise ValueError(f"不支援的處理模式：{processing_mode}")
    items = collect_source_items(input_paths, allowed_extensions, excluded_paths=[output_dir])
    plan = create_scan_plan_from_items(
        items,
        output_dir,
        manifest,
        args.force,
        output_format,
        model=args.model,
        base_url=args.base_url,
        processing_mode=processing_mode,
    )
    if plan.manifest_changed:
        manifest.save()

    api_key = getattr(args, "api_key", None) or os.environ.get("DEEPSEEK_API_KEY")
    vision_factory = VisionClientFactory(api_key, args.base_url, args.model)
    converted = 0
    failed = 0
    image_requests = 0
    enhancement_requests = 0

    if plan.jobs:
        with ThreadPoolExecutor(max_workers=args.workers, thread_name_prefix="doc-to-md") as executor:
            futures = {
                executor.submit(
                    convert_job,
                    job,
                    args.model,
                    output_format,
                    processing_mode,
                    vision_factory,
                ): job
                for job in plan.jobs
            }
            with tqdm(total=len(futures), desc="轉換文件", unit="file", disable=quiet) as progress:
                for future in as_completed(futures):
                    result = future.result()
                    image_requests += result.image_requests
                    enhancement_requests += result.enhancement_requests
                    if result.success:
                        converted += 1
                        for warning in result.warnings:
                            logger.warning("%s\t%s", result.job.relative_source.as_posix(), warning)
                    else:
                        failed += 1
                        logger.error("%s\t%s", result.job.relative_source.as_posix(), result.error)
                    update_manifest(
                        manifest,
                        result,
                        args.model,
                        base_url=args.base_url,
                        output_format=output_format,
                        processing_mode=processing_mode,
                    )
                    manifest.save()
                    progress.update(1)
                    if progress_callback:
                        progress_callback(result, converted + failed, len(futures))

    summary = {
        "total": plan.total,
        "skipped": plan.skipped,
        "converted": converted,
        "failed": failed,
        "image_requests": image_requests,
        "enhancement_requests": enhancement_requests,
    }
    setattr(args, "run_summary", summary)
    if not quiet:
        print("\n執行摘要")
        print(f"  掃描總檔案數: {plan.total}")
        print(f"  略過未變更數: {plan.skipped}")
        print(f"  成功轉換數: {converted}")
        print(f"  失敗檔案數: {failed}")
        print(f"  VLM 圖片請求數: {image_requests}")
        print(f"  AI 整理請求數: {enhancement_requests}")
        if failed:
            print(f"  錯誤記錄: {output_dir / ERROR_LOG_FILENAME}")
    close_logger(logger)
    return 1 if failed else 0


def run_pipeline(args: argparse.Namespace) -> int:
    input_dir, output_dir = validate_directories(Path(args.input_dir), Path(args.output_dir))
    args.output_dir = str(output_dir)
    return run_pipeline_paths(args, [input_dir])


def positive_int(value: str) -> int:
    parsed = int(value)
    if parsed < 1:
        raise argparse.ArgumentTypeError("必須是大於 0 的整數")
    return parsed


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="批次將企業文件轉換為 AnythingLLM Markdown")
    parser.add_argument("--input-dir", required=True, help="來源文件根目錄")
    parser.add_argument("--output-dir", required=True, help="Markdown 輸出根目錄")
    parser.add_argument("--workers", type=positive_int, default=4, help="並行處理數（預設：4）")
    parser.add_argument("--force", action="store_true", help="忽略 Manifest，強制重新轉換所有文件")
    parser.add_argument(
        "--output-format",
        choices=("md", "txt", "json"),
        default="md",
        help="輸出格式（預設：md）",
    )
    parser.add_argument(
        "--processing-mode",
        choices=("hybrid", "ai-enhanced"),
        default="hybrid",
        help="處理模式：hybrid 或轉換後再由 AI 整理的 ai-enhanced（預設：hybrid）",
    )
    parser.add_argument("--model", default=DEFAULT_MODEL, help=f"DeepSeek AI 模型（預設：{DEFAULT_MODEL}）")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help=f"DeepSeek API Base URL（預設：{DEFAULT_BASE_URL}）")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    try:
        from dotenv import load_dotenv

        load_dotenv()
        args = build_parser().parse_args(argv)
        return run_pipeline(args)
    except KeyboardInterrupt:
        print("\n已由使用者中止。", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"錯誤：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
