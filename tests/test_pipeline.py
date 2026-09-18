import argparse
import os
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import doc_to_md_pipeline as pipeline


class IncrementalScanTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.input_dir = root / "input"
        self.output_dir = root / "output"
        self.input_dir.mkdir()
        self.output_dir.mkdir()
        self.source = self.input_dir / "nested" / "sample.txt"
        self.source.parent.mkdir()
        self.source.write_text("first", encoding="utf-8")
        self.manifest = pipeline.Manifest(self.output_dir / pipeline.MANIFEST_FILENAME)

    def tearDown(self):
        self.temp.cleanup()

    def mark_success(self, job):
        job.output.parent.mkdir(parents=True, exist_ok=True)
        job.output.write_text("converted", encoding="utf-8")
        result = pipeline.ConversionResult(job=job, success=True)
        pipeline.update_manifest(self.manifest, result, pipeline.DEFAULT_MODEL)

    def test_first_scan_queues_file(self):
        plan = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.assertEqual(plan.total, 1)
        self.assertEqual(plan.skipped, 0)
        self.assertEqual(len(plan.jobs), 1)
        self.assertEqual(plan.jobs[0].output_relative.as_posix(), "nested/sample.md")

    def test_unchanged_mtime_fast_skips(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        second = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.assertEqual(second.skipped, 1)
        self.assertFalse(second.jobs)

    def test_mtime_only_change_hash_skips_and_updates_manifest(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        old_mtime = self.source.stat().st_mtime
        os.utime(self.source, (old_mtime + 5, old_mtime + 5))
        second = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.assertEqual(second.skipped, 1)
        self.assertFalse(second.jobs)
        self.assertTrue(second.manifest_changed)
        self.assertEqual(
            self.manifest.entries["nested/sample.txt"]["last_modified"], self.source.stat().st_mtime
        )

    def test_content_change_queues_file(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        time.sleep(0.01)
        self.source.write_text("second", encoding="utf-8")
        second = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.assertEqual(len(second.jobs), 1)
        self.assertNotEqual(second.jobs[0].sha256, first.jobs[0].sha256)

    def test_missing_output_queues_file(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        first.jobs[0].output.unlink()
        second = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.assertEqual(len(second.jobs), 1)

    def test_force_queues_unchanged_file(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        second = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, True)
        self.assertEqual(len(second.jobs), 1)
        self.assertEqual(second.skipped, 0)

    def test_model_change_queues_unchanged_file(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        items = [pipeline.SourceItem(self.source, self.source.relative_to(self.input_dir))]
        second = pipeline.create_scan_plan_from_items(
            items,
            self.output_dir,
            self.manifest,
            False,
            model="different-vision-model",
        )
        self.assertEqual(len(second.jobs), 1)

    def test_processing_mode_change_queues_unchanged_file(self):
        first = pipeline.create_scan_plan(self.input_dir, self.output_dir, self.manifest, False)
        self.mark_success(first.jobs[0])
        items = [pipeline.SourceItem(self.source, self.source.relative_to(self.input_dir))]
        second = pipeline.create_scan_plan_from_items(
            items,
            self.output_dir,
            self.manifest,
            False,
            processing_mode="ai-enhanced",
        )
        self.assertEqual(len(second.jobs), 1)


class MappingTests(unittest.TestCase):
    def test_same_stem_uses_extension_to_avoid_collision(self):
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            pdf = input_dir / "report.pdf"
            docx = input_dir / "report.docx"
            pdf.touch()
            docx.touch()
            mapping = pipeline.build_output_mapping([pdf, docx], input_dir, output_dir)
            self.assertEqual(mapping[pdf][1].as_posix(), "report.pdf.md")
            self.assertEqual(mapping[docx][1].as_posix(), "report.docx.md")


class EndToEndTextTests(unittest.TestCase):
    def test_text_conversion_and_second_run_skip(self):
        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "hello.txt").write_text("Hello Markdown", encoding="utf-8")
            args = argparse.Namespace(
                input_dir=str(input_dir),
                output_dir=str(output_dir),
                workers=2,
                force=False,
                model=pipeline.DEFAULT_MODEL,
                base_url=pipeline.DEFAULT_BASE_URL,
            )
            self.assertEqual(pipeline.run_pipeline(args), 0)
            markdown = (output_dir / "hello.md").read_text(encoding="utf-8")
            self.assertIn('original_file: "hello.txt"', markdown)
            self.assertIn("Hello Markdown", markdown)
            self.assertEqual(pipeline.run_pipeline(args), 0)
            manifest = pipeline.Manifest(output_dir / pipeline.MANIFEST_FILENAME)
            manifest.load()
            self.assertEqual(manifest.entries["hello.txt"]["status"], "success")


class ModernFormatParserTests(unittest.TestCase):
    def test_docx_text_and_table(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "sample.docx"
            document = Document()
            document.add_heading("Architecture", level=1)
            document.add_paragraph("Controller details")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Clock"
            table.cell(1, 1).text = "100 MHz"
            document.save(path)
            result = pipeline.extract_docx(path)
            text = "\n".join(str(part.content) for part in result.parts)
            self.assertIn("# Architecture", text)
            self.assertIn("Controller details", text)
            self.assertIn("100 MHz", text)

    def test_xlsx_all_sheets(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "sample.xlsx"
            workbook = Workbook()
            workbook.active.title = "Overview"
            workbook.active.append(["Part", "Count"])
            workbook.active.append(["FPGA", 2])
            workbook.create_sheet("Empty")
            workbook.save(path)
            result = pipeline.extract_xlsx(path)
            text = "\n".join(str(part.content) for part in result.parts)
            self.assertIn("## 工作表：Overview", text)
            self.assertIn("FPGA", text)
            self.assertIn("## 工作表：Empty", text)

    def test_pptx_text_without_libreoffice(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "System Overview"
            presentation.save(path)
            result = pipeline.extract_pptx(path)
            text = "\n".join(str(part.content) for part in result.parts)
            self.assertIn("System Overview", text)

    def test_pdf_native_text(self):
        import pymupdf as fitz

        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "sample.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Native PDF text")
            document.save(path)
            document.close()
            result = pipeline.extract_pdf(path)
            text = "\n".join(str(part.content) for part in result.parts if part.kind == "text")
            self.assertIn("Native PDF text", text)


@unittest.skipUnless(pipeline.find_soffice(), "LibreOffice is required for legacy Office tests")
class LegacyOfficeIntegrationTests(unittest.TestCase):
    def test_legacy_doc_round_trip(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            modern = root / "legacy-source.docx"
            document = Document()
            document.add_paragraph("Legacy DOC marker")
            document.save(modern)
            legacy = pipeline.libreoffice_convert(modern, ".doc", root / "legacy")
            result = pipeline.extract_document(legacy)
            text = "\n".join(str(part.content) for part in result.parts if part.kind == "text")
            self.assertIn("Legacy DOC marker", text)

    def test_legacy_xls_round_trip(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            modern = root / "legacy-source.xlsx"
            workbook = Workbook()
            workbook.active.append(["Legacy XLS marker", 42])
            workbook.save(modern)
            legacy = pipeline.libreoffice_convert(modern, ".xls", root / "legacy")
            result = pipeline.extract_document(legacy)
            text = "\n".join(str(part.content) for part in result.parts if part.kind == "text")
            self.assertIn("Legacy XLS marker", text)

    def test_legacy_ppt_round_trip(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as name:
            root = Path(name)
            modern = root / "legacy-source.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "Legacy PPT marker"
            presentation.save(modern)
            legacy = pipeline.libreoffice_convert(modern, ".ppt", root / "legacy")
            result = pipeline.extract_document(legacy)
            text = "\n".join(str(part.content) for part in result.parts if part.kind == "text")
            self.assertIn("Legacy PPT marker", text)


class OutputAndApiTests(unittest.TestCase):
    def test_json_output_contains_markdown_body(self):
        markdown = '---\noriginal_file: "a.txt"\n---\n\n# Hello\n'
        rendered = pipeline.serialize_output(
            markdown, "json", Path("a.txt"), pipeline.DEFAULT_MODEL, "2026-01-01T00:00:00Z"
        )
        self.assertIn('"content_markdown": "# Hello"', rendered)
        self.assertIn('"processing_mode": "hybrid"', rendered)

    def test_model_list_is_sorted_and_deduplicated(self):
        fake_response = SimpleNamespace(
            data=[SimpleNamespace(id="model-z"), SimpleNamespace(id="model-a"), SimpleNamespace(id="model-a")]
        )
        fake_client = SimpleNamespace(models=SimpleNamespace(list=lambda: fake_response))
        with patch("openai.OpenAI", return_value=fake_client):
            models = pipeline.list_available_models("secret", "https://example.test/v1")
        self.assertEqual(models, ["model-a", "model-z"])


class AiEnhancedTests(unittest.TestCase):
    def test_tables_are_protected_from_ai_enhancement(self):
        body = "# Weekly report\n\nNarrative text.\n\n| Name | Value |\n| --- | --- |\n| Clock | 100 MHz |"
        segments = pipeline.split_markdown_for_enhancement(body, max_chars=1000)
        self.assertTrue(any(eligible for _, eligible in segments))
        protected = [content for content, eligible in segments if not eligible]
        self.assertEqual(len(protected), 1)
        self.assertIn("100 MHz", protected[0])

    def test_sheet_heading_next_to_table_is_also_protected(self):
        body = "## 工作表：Overview\n\n| Part | Count |\n| --- | --- |\n| FPGA | 2 |"
        segments = pipeline.split_markdown_for_enhancement(body, max_chars=1000)
        self.assertTrue(segments)
        self.assertFalse(any(eligible for _, eligible in segments))

    def test_fenced_code_with_blank_lines_is_protected(self):
        body = "```python\nvalue = 12\n\nprint(value)\n```"
        segments = pipeline.split_markdown_for_enhancement(body, max_chars=1000)
        self.assertEqual(segments, [(body, False)])

    def test_numeric_change_fails_safety_validation(self):
        valid, reason = pipeline.validate_enhanced_markdown(
            "Revenue was 1200 in 2026.", "## Revenue\n\nRevenue was 1250 in 2026."
        )
        self.assertFalse(valid)
        self.assertIn("不一致", reason)

    def test_rewritten_text_fails_safety_validation(self):
        source = "The controller validates each document and preserves the original technical content."
        enhanced = "The assistant invents a completely different summary with unrelated operational guidance."
        valid, reason = pipeline.validate_enhanced_markdown(source, enhanced)
        self.assertFalse(valid)
        self.assertIn("文字內容", reason)

    def test_document_enhancement_preserves_frontmatter_and_tables(self):
        class FakeClient:
            def enhance_markdown(self, markdown, chunk_index, chunk_count):
                return markdown.replace("# Weekly report", "# Weekly Report")

        factory = SimpleNamespace(get=lambda: FakeClient())
        source = (
            '---\noriginal_file: "report.md"\n---\n\n'
            "# Weekly report\n\nCompleted in 2026.\n\n"
            "| Name | Value |\n| --- | --- |\n| Clock | 100 MHz |\n"
        )
        enhanced, requests, warnings = pipeline.enhance_markdown_document(source, factory)
        self.assertTrue(enhanced.startswith('---\noriginal_file: "report.md"\n---\n\n'))
        self.assertIn("# Weekly Report", enhanced)
        self.assertIn("| Clock | 100 MHz |", enhanced)
        self.assertEqual(requests, 1)
        self.assertFalse(warnings)

    def test_unsafe_ai_result_falls_back_to_original(self):
        class UnsafeClient:
            def enhance_markdown(self, markdown, chunk_index, chunk_count):
                return markdown.replace("2026", "2027")

        factory = SimpleNamespace(get=lambda: UnsafeClient())
        source = '---\noriginal_file: "report.md"\n---\n\n# Report\n\nCompleted in 2026.\n'
        enhanced, requests, warnings = pipeline.enhance_markdown_document(source, factory)
        self.assertIn("2026", enhanced)
        self.assertNotIn("2027", enhanced)
        self.assertEqual(requests, 1)
        self.assertEqual(len(warnings), 1)


if __name__ == "__main__":
    unittest.main()
